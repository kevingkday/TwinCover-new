import io
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class ExportService:
    """Exports generated insurance renewal letters to DOCX and PDF formats."""

    @staticmethod
    def export_to_docx(
        title: str,
        content: str,
        client_name: Optional[str] = None,
        policy_number: Optional[str] = None
    ) -> bytes:
        """Create a polished commercial broker letter in DOCX format."""
        doc = Document()

        # Set standard 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Header Agency Banner
        header_para = doc.add_paragraph()
        header_run = header_para.add_run("TWINCOVER COMMERCIAL ADVISORY GROUP\n")
        header_run.bold = True
        header_run.font.size = Pt(13)
        header_run.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900

        sub_run = header_para.add_run("Commercial Lines Risk Management & Policy Summaries")
        sub_run.font.size = Pt(9.5)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)  # Slate 500

        # Title
        title_para = doc.add_paragraph()
        title_para.space_before = Pt(14)
        title_para.space_after = Pt(10)
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(30, 41, 59)

        # Parse markdown lines into Word elements
        lines = content.splitlines()
        in_table = False
        table_rows = []

        for line in lines:
            stripped = line.strip()

            # Handle Markdown table lines
            if stripped.startswith("|") and stripped.endswith("|"):
                if not in_table:
                    in_table = True
                    table_rows = []
                # Check if it's separator row
                if re.match(r"^\|[\s\-:|]+\|$", stripped):
                    continue
                cells = [c.strip() for c in stripped.split("|")[1:-1]]
                table_rows.append(cells)
                continue
            else:
                if in_table and table_rows:
                    # Render docx table
                    num_cols = max(len(r) for r in table_rows)
                    table = doc.add_table(rows=len(table_rows), cols=num_cols)
                    table.style = 'Light Shading Accent 1' if 'Light Shading Accent 1' in [s.name for s in doc.styles] else 'Table Grid'
                    for r_idx, row_data in enumerate(table_rows):
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx < num_cols:
                                cell = table.cell(r_idx, c_idx)
                                cell.text = cell_value
                                if r_idx == 0:
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                    in_table = False
                    table_rows = []

            if not stripped:
                continue

            # Heading 1
            if stripped.startswith("# "):
                h = doc.add_paragraph()
                h.space_before = Pt(12)
                h.space_after = Pt(4)
                r = h.add_run(stripped[2:])
                r.bold = True
                r.font.size = Pt(14)
                r.font.color.rgb = RGBColor(30, 58, 138)  # Blue 900
            # Heading 2
            elif stripped.startswith("## "):
                h = doc.add_paragraph()
                h.space_before = Pt(10)
                h.space_after = Pt(4)
                r = h.add_run(stripped[3:])
                r.bold = True
                r.font.size = Pt(12.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            # Heading 3
            elif stripped.startswith("### "):
                h = doc.add_paragraph()
                h.space_before = Pt(8)
                h.space_after = Pt(2)
                r = h.add_run(stripped[4:])
                r.bold = True
                r.font.size = Pt(11)
            # Bullet list
            elif stripped.startswith("- ") or stripped.startswith("* "):
                p = doc.add_paragraph(style='List Bullet')
                p.space_after = Pt(3)
                ExportService._format_styled_text(p, stripped[2:])
            # Blockquote
            elif stripped.startswith("> "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.space_before = Pt(4)
                p.space_after = Pt(4)
                r = p.add_run(stripped[2:])
                r.italic = True
                r.font.color.rgb = RGBColor(71, 85, 105)
            # Standard paragraph
            else:
                p = doc.add_paragraph()
                p.space_after = Pt(5)
                ExportService._format_styled_text(p, stripped)

        # Flush any trailing table
        if in_table and table_rows:
            num_cols = max(len(r) for r in table_rows)
            table = doc.add_table(rows=len(table_rows), cols=num_cols)
            for r_idx, row_data in enumerate(table_rows):
                for c_idx, cell_value in enumerate(row_data):
                    if c_idx < num_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_value

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _format_styled_text(paragraph, text: str):
        """Parse **bold** and *italic* tokens into Word runs."""
        tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                r = paragraph.add_run(token[2:-2])
                r.bold = True
            elif token.startswith("*") and token.endswith("*"):
                r = paragraph.add_run(token[1:-1])
                r.italic = True
            else:
                paragraph.add_run(token)

    @staticmethod
    def export_to_pdf(
        title: str,
        content: str,
        client_name: Optional[str] = None,
        policy_number: Optional[str] = None
    ) -> bytes:
        """Create a print-ready commercial letter in PDF format via ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )
        styles = getSampleStyleSheet()

        # Custom professional styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=12
        )
        h2_style = ParagraphStyle(
            'DocH2',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1e3a8a"),
            spaceBefore=10,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6
        )
        bullet_style = ParagraphStyle(
            'DocBullet',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155"),
            leftIndent=15,
            spaceAfter=4
        )

        story = []
        # Header banner
        story.append(Paragraph("<b>TWINCOVER COMMERCIAL ADVISORY GROUP</b>", title_style))
        story.append(Paragraph("<font size=9 color='#64748b'>Commercial Lines Insurance Placement & Policy Renewal Analysis</font>", body_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14))

        # Title
        story.append(Paragraph(f"<b>{title}</b>", title_style))
        story.append(Spacer(1, 6))

        for line in content.splitlines():
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 4))
                continue

            # Convert markdown styling to simple HTML tags for ReportLab
            html_line = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", stripped)
            html_line = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", html_line)

            if stripped.startswith("# "):
                story.append(Paragraph(html_line[2:], title_style))
            elif stripped.startswith("## "):
                story.append(Paragraph(html_line[3:], h2_style))
            elif stripped.startswith("### "):
                story.append(Paragraph(f"<b>{html_line[4:]}</b>", h2_style))
            elif stripped.startswith("- ") or stripped.startswith("* "):
                story.append(Paragraph(f"&bull; {html_line[2:]}", bullet_style))
            else:
                story.append(Paragraph(html_line, body_style))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def render_docx_template(
        template_path: str,
        fields: dict,
        broker_info: dict
    ) -> bytes:
        """Merge verified policy fields and dynamic coverages into an uploaded Word (.docx) template."""
        import json
        from datetime import datetime
        from docxtpl import DocxTemplate

        doc = DocxTemplate(template_path)

        # Parse coverages list
        coverages_list = fields.get("coverages") or []
        if isinstance(coverages_list, str):
            try:
                coverages_list = json.loads(coverages_list)
            except Exception:
                coverages_list = []

        if coverages_list and isinstance(coverages_list, list):
            lines = []
            for c in coverages_list:
                name = c.get("name") or "Standard Coverage"
                limit = c.get("limit") or "Included"
                ded = c.get("deductible") or "Standard / Nil"
                lines.append(f"• {name}: Limit {limit} | Deductible: {ded}")
            coverages_table_text = "\n".join(lines)
        else:
            coverages_table_text = f"• Commercial General Liability: {fields.get('cgl_limit') or '$2,000,000'}\n• Cyber Deductible: {fields.get('cyber_deductible') or '$5,000'}"

        context = {
            "insured_name": fields.get("insured_name") or "[Client Insured Name]",
            "policy_number": fields.get("policy_number") or "[Policy # Pending]",
            "insurer": fields.get("insurer") or "[Insurance Carrier]",
            "effective_date": fields.get("effective_date") or "[Effective Date]",
            "expiry_date": fields.get("expiry_date") or "[Expiry Date]",
            "total_premium": fields.get("total_premium") or "$0.00",
            "cgl_limit": fields.get("cgl_limit") or "$2,000,000",
            "property_limit": fields.get("property_limit") or "$10,000,000",
            "cyber_limit": fields.get("cyber_limit") or "$2,000,000",
            "cyber_deductible": fields.get("cyber_deductible") or "$5,000",
            "coverages": coverages_list,
            "coverages_table": coverages_table_text,
            "broker_name": broker_info.get("full_name") or "Commercial Insurance Account Executive",
            "agency_name": broker_info.get("agency_name") or "TwinCover Commercial Advisory Group",
            "broker_email": broker_info.get("email") or "broker@twincover.com",
            "current_date": broker_info.get("current_date") or datetime.now().strftime("%B %d, %Y")
        }

        doc.render(context)
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()


export_service = ExportService()
